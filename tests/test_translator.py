from __future__ import annotations

import pytest
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from translator.cli import main
from translator.pdf_numbers import find_step_numbers
from translator.pdf_steps import SectionNotFoundError, extract_step_names

STYLES = getSampleStyleSheet()
GRID = TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)])


def _table(rows):
    t = Table(rows, repeatRows=1)
    t.setStyle(GRID)
    return t


def make_pdf1(path, with_bookmarks: bool, many_rows: bool = False):
    step_rows = [["Step", "Description"], ["T00", "Init"], ["T01", "Warm up"],
                 ["S01", "First"], ["S02", "Second"]]
    if many_rows:
        step_rows += [[f"S{i:02d}", f"Step {i}"] for i in range(3, 60)]

    story = [
        Paragraph("Table of contents", STYLES["Heading1"]),
        Paragraph("2 Overview ........ 2", STYLES["Normal"]),
        Paragraph("3.1 Step list ........ 3", STYLES["Normal"]),
        Paragraph("3.2 Other steps ........ 4", STYLES["Normal"]),
        PageBreak(),
        Paragraph("2 Overview", STYLES["Heading1"]),
        _table([["Item", "Value"], ["X01", "wrong table"]]),
        PageBreak(),
        Paragraph("3.1 Step list", STYLES["Heading2"]),
        Spacer(1, 12),
        _table(step_rows),
        Paragraph("3.2 Other steps", STYLES["Heading2"]),
        _table([["Step", "Description"], ["Z99", "not in 3.1"]]),
    ]

    doc = SimpleDocTemplate(str(path), pagesize=A4)
    if with_bookmarks:
        class Doc(SimpleDocTemplate):
            def afterFlowable(self, flowable):
                if isinstance(flowable, Paragraph) and flowable.style.name.startswith("Heading"):
                    text = flowable.getPlainText()
                    if text[0].isdigit():
                        key = text.replace(" ", "_")
                        self.canv.bookmarkPage(key)
                        self.canv.addOutlineEntry(text, key, 0)
        doc = Doc(str(path), pagesize=A4)
    doc.build(story)


def make_pdf2(path):
    pages = [
        ["Step properties", "Name: T00", "number: 100", "type: start"],
        ["Name: T01", "number: 101"],
        ["Name: S01_1", "number: 201", "previous: T01"],
        ["Name: S01_2", "number:", "202"],  # value on the next line
        ["Name: S01_3", "Number: 203"],
        ["Name: S010", "number: 999"],  # must not be taken for S01
        ["Name: S02", "number: 300"],
        ["Name: Z99", "number: 777"],
    ]
    c = canvas.Canvas(str(path), pagesize=A4)
    for lines in pages:
        y = 800
        for line in lines:
            c.drawString(72, y, line)
            y -= 20
        c.showPage()
    c.save()


@pytest.fixture
def pdfs(tmp_path):
    p1, p2 = tmp_path / "steps.pdf", tmp_path / "numbers.pdf"
    make_pdf1(p1, with_bookmarks=False)
    make_pdf2(p2)
    return p1, p2


@pytest.mark.parametrize("with_bookmarks", [False, True])
def test_extract_steps_from_section(tmp_path, with_bookmarks):
    p1 = tmp_path / "steps.pdf"
    make_pdf1(p1, with_bookmarks=with_bookmarks)
    assert extract_step_names(str(p1), "3.1") == ["T00", "T01", "S01", "S02"]
    assert extract_step_names(str(p1), "3.2") == ["Z99"]


def test_table_spanning_pages(tmp_path):
    p1 = tmp_path / "steps.pdf"
    make_pdf1(p1, with_bookmarks=False, many_rows=True)
    steps = extract_step_names(str(p1), "3.1")
    assert steps[:4] == ["T00", "T01", "S01", "S02"]
    assert steps[-1] == "S59"
    assert "Z99" not in steps


def test_missing_section(pdfs):
    with pytest.raises(SectionNotFoundError):
        extract_step_names(str(pdfs[0]), "7.4")


def test_numbers_with_variants(pdfs):
    res = find_step_numbers(str(pdfs[1]), ["T00", "T01", "S01", "S02", "S03"])
    assert res["T00"].numbers == ["100"]
    assert res["T01"].numbers == ["101"]
    assert res["S01"].numbers == ["201", "202", "203"]
    assert [m.variant for m in res["S01"].matches] == ["S01_1", "S01_2", "S01_3"]
    assert res["S02"].numbers == ["300"]
    assert res["S03"].numbers == []


def test_name_label(pdfs):
    res = find_step_numbers(str(pdfs[1]), ["T01", "S01"], name_label="Name")
    assert res["T01"].numbers == ["101"]
    assert res["S01"].numbers == ["201", "202", "203"]


def test_end_to_end_excel(pdfs, tmp_path):
    out = tmp_path / "mapping.xlsx"
    rc = main(["--pdf1", str(pdfs[0]), "--section", "3.1", "--pdf2", str(pdfs[1]),
               "--output", str(out)])
    assert rc == 0
    ws = load_workbook(out).active
    rows = [tuple(c.value for c in r) for r in ws.iter_rows(min_row=1, max_col=2)]
    assert rows == [
        ("Step", "Number"),
        ("T00", 100),
        ("T01", 101),
        ("S01", "201 OR 202 OR 203"),
        ("S02", 300),
    ]


def test_custom_layout(pdfs, tmp_path):
    out = tmp_path / "mapping.xlsx"
    main(["--pdf1", str(pdfs[0]), "--section", "3.1", "--pdf2", str(pdfs[1]),
          "--output", str(out), "--step-column", "C", "--number-column", "E",
          "--start-row", "3", "--step-header", "Name", "--number-header", "No."])
    ws = load_workbook(out).active
    assert ws["C3"].value == "Name" and ws["E3"].value == "No."
    assert ws["C6"].value == "S01" and ws["E6"].value == "201 OR 202 OR 203"
