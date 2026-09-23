"""Write the step -> number mapping to a new Excel file."""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string

from .pdf_numbers import StepResult

NOT_FOUND_FILL = PatternFill("solid", start_color="FFFF00", end_color="FFFF00")


def _cell_value(numbers: list[str], separator: str):
    if len(numbers) == 1 and numbers[0].isdigit() and not numbers[0].startswith("0"):
        return int(numbers[0])
    return f" {separator} ".join(numbers) if numbers else None


def step_digits(step: str) -> int | None:
    """The number inside a step name: "T10" -> 10, "S02" -> 2, "T00" -> 0."""
    m = re.search(r"\d+", step)
    return int(m.group()) if m else None


def natural_key(step: str):
    """Alphabetical order that compares the digits as numbers: S2 < S02b < S10."""
    return [int(part) if part.isdigit() else part.upper() for part in re.split(r"(\d+)", step)]


@dataclass
class Sheet:
    name: str
    results: list[StepResult] = field(default_factory=list)
    error: str | None = None  # shown on the sheet when the sequence could not be read


def write_mapping(output_path: str, results: list[StepResult], *, sheet_name: str = "Mapping",
                  checks: list[str] | None = None, **layout) -> None:
    """Create `output_path` with a single sheet (see write_workbook for `layout`)."""
    write_workbook(output_path, [Sheet(sheet_name, results)], checks=checks, **layout)


def write_workbook(
    output_path: str,
    sheets: list[Sheet],
    *,
    checks: list[str] | None = None,
    step_column: str = "A",
    step_number_column: str | None = "B",
    number_column: str = "C",
    pages_column: str | None = None,
    start_row: int = 1,
    step_header: str | None = "Step",
    step_number_header: str | None = "Step no.",
    number_header: str | None = "Number",
    pages_header: str = "PDF 2 pages",
    separator: str = "OR",
    not_found_text: str = "",
    sort: bool = True,
) -> None:
    """Create `output_path` with one sheet per sequence, one row per step
    (alphabetical when `sort`).

    `step_number_column` receives the digits of the step name (T10 -> 10).
    Headers are written on `start_row` (unless both headers are None) and the data
    starts on the row below. Steps without a number are highlighted in yellow.
    `checks` (points for a human to verify) are listed on a separate "Check" sheet.
    """
    wb = Workbook()
    wb.remove(wb.active)

    step_col = column_index_from_string(step_column.upper())
    step_no_col = column_index_from_string(step_number_column.upper()) if step_number_column else None
    number_col = column_index_from_string(number_column.upper())
    pages_col = column_index_from_string(pages_column.upper()) if pages_column else None

    for sheet in sheets:
        ws = wb.create_sheet(sheet.name[:31])
        row = start_row
        if step_header is not None or number_header is not None:
            headers = [(step_col, step_header), (number_col, number_header)]
            if step_no_col:
                headers.append((step_no_col, step_number_header))
            if pages_col:
                headers.append((pages_col, pages_header))
            for col, text in headers:
                if text is not None:
                    ws.cell(row=row, column=col, value=text).font = Font(bold=True)
            row += 1

        if sheet.error:
            cell = ws.cell(row=row, column=step_col, value=f"ERROR: {sheet.error}")
            cell.font = Font(bold=True, color="FF0000")
            ws.sheet_properties.tabColor = "FF0000"

        results = sorted(sheet.results, key=lambda r: natural_key(r.step)) if sort else sheet.results
        for result in results:
            ws.cell(row=row, column=step_col, value=result.step)
            if step_no_col:
                ws.cell(row=row, column=step_no_col, value=step_digits(result.step))
            value = _cell_value(result.numbers, separator)
            number_cell = ws.cell(row=row, column=number_col, value=value)
            if value is None:
                if not_found_text:
                    number_cell.value = not_found_text
                number_cell.fill = NOT_FOUND_FILL
            if pages_col and result.pages:
                ws.cell(row=row, column=pages_col, value=", ".join(map(str, result.pages)))
            row += 1

        for col, width in ((step_col, 15), (step_no_col, 10), (number_col, 40), (pages_col, 15)):
            if col:
                ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    if checks:
        check = wb.create_sheet("Check")
        check.cell(row=1, column=1, value="Please verify").font = Font(bold=True)
        for i, text in enumerate(checks, start=2):
            check.cell(row=i, column=1, value=text)
        check.column_dimensions["A"].width = 110
        check.sheet_properties.tabColor = "FF0000"

    wb.save(output_path)
